from __future__ import annotations

import json
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urljoin

from openpyxl import load_workbook
from selenium import webdriver
from selenium.webdriver.chrome.options import Options

from action_keywords import ActionKeywords

import sys
sys.path.append(str(Path(__file__).resolve().parents[1]))
from report_utils import write_html_report, write_json_report  # noqa: E402


SUPPORTED_KEYWORDS = {
    "open",
    "configureApi",
    "waitVisible",
    "type",
    "click",
    "assertValue",
    "assertTextContains",
    "assertUrlContains",
    "assertLocalStorageObjectNotEmpty",
}
KEYWORDS_REQUIRING_OBJECT = {"open", "waitVisible", "type", "click", "assertValue", "assertLocalStorageObjectNotEmpty"}


def clean(value: Any) -> str:
    return str(value or "").strip()


def has_data(row: dict[str, Any]) -> bool:
    return any(clean(value) for value in row.values())


def read_steps(test_data_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(test_data_path, data_only=True)
    sheet_name = "TestSteps" if "TestSteps" in workbook.sheetnames else "WebsiteKeywords"
    if sheet_name not in workbook.sheetnames:
        raise RuntimeError(f'Missing sheet "TestSteps" in {test_data_path}')

    sheet = workbook[sheet_name]
    headers = [clean(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    steps: list[dict[str, Any]] = []
    for excel_row in sheet.iter_rows(min_row=2, values_only=True):
        row = {headers[index]: excel_row[index] if index < len(excel_row) else "" for index in range(len(headers))}
        if not has_data(row):
            continue
        steps.append(
            {
                "testCaseId": clean(row.get("testCaseId") or row.get("id")),
                "testCaseName": clean(row.get("testCaseName") or row.get("name")),
                "stepOrder": int(row.get("stepOrder") or 0),
                "keyword": clean(row.get("keyword")),
                "objectName": clean(row.get("objectName") or row.get("target")),
                "testData": clean(row.get("testData") or row.get("value")),
                "description": clean(row.get("description")),
                "runFlag": clean(row.get("runFlag") or "Y").upper(),
            }
        )
    return [step for step in steps if step["runFlag"] != "N"]


def validate_steps(steps: list[dict[str, Any]]) -> None:
    for step in steps:
        if not step["testCaseId"]:
            raise RuntimeError("Keyword Excel has a row without testCaseId.")
        if not step["keyword"]:
            raise RuntimeError(f"Test case {step['testCaseId']} has a row without keyword.")
        if step["keyword"] not in SUPPORTED_KEYWORDS:
            raise RuntimeError(f"Test case {step['testCaseId']} uses unsupported keyword \"{step['keyword']}\".")
        if step["keyword"] in KEYWORDS_REQUIRING_OBJECT and not step["objectName"]:
            raise RuntimeError(f"Test case {step['testCaseId']}, keyword {step['keyword']} requires objectName.")


def group_steps(steps: list[dict[str, Any]]) -> list[dict[str, Any]]:
    cases: dict[str, dict[str, Any]] = {}
    for step in steps:
        test_case = cases.setdefault(
            step["testCaseId"],
            {"id": step["testCaseId"], "name": step["testCaseName"] or step["testCaseId"], "steps": []},
        )
        test_case["steps"].append(step)
    return [
        {**test_case, "steps": sorted(test_case["steps"], key=lambda item: item["stepOrder"])}
        for test_case in cases.values()
    ]


def build_driver() -> webdriver.Chrome:
    options = Options()
    options.add_argument("--headless=new")
    options.add_argument("--disable-gpu")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--ignore-certificate-errors")
    options.add_argument("--window-size=1366,768")
    return webdriver.Chrome(options=options)


def run_case(driver, actions: ActionKeywords, frontend_base_url: str, test_case: dict[str, Any]) -> dict[str, Any]:
    started_at = time.time()
    step_results = []
    try:
        driver.get(urljoin(frontend_base_url.rstrip("/") + "/", "/index.html".lstrip("/")))
        driver.execute_script("localStorage.clear(); sessionStorage.clear();")
        for step in test_case["steps"]:
            step_started_at = time.time()
            detail = actions.execute(step)
            step_results.append(
                {
                    "stepOrder": step["stepOrder"],
                    "keyword": step["keyword"],
                    "objectName": step["objectName"],
                    "testData": step["testData"],
                    "description": step["description"],
                    "status": "passed",
                    "detail": detail,
                    "timeMs": int((time.time() - step_started_at) * 1000),
                }
            )
        return {
            "id": test_case["id"],
            "name": test_case["name"],
            "status": "passed",
            "steps": step_results,
            "timeMs": int((time.time() - started_at) * 1000),
        }
    except Exception as error:
        failed_step = test_case["steps"][len(step_results)] if len(step_results) < len(test_case["steps"]) else {}
        step_results.append(
            {
                "stepOrder": failed_step.get("stepOrder"),
                "keyword": failed_step.get("keyword"),
                "objectName": failed_step.get("objectName"),
                "testData": failed_step.get("testData"),
                "description": failed_step.get("description"),
                "status": "failed",
                "error": str(error),
            }
        )
        return {
            "id": test_case["id"],
            "name": test_case["name"],
            "status": "failed",
            "steps": step_results,
            "error": str(error),
            "timeMs": int((time.time() - started_at) * 1000),
        }


def run_keyword_driven_tests(options: dict[str, Any]) -> dict[str, Any]:
    reports_dir = Path(options["reportsDir"])
    reports_dir.mkdir(parents=True, exist_ok=True)
    test_data_path = Path(options["testDataPath"])
    object_repository_path = Path(__file__).resolve().parent / "object-repository.json"
    object_repository = json.loads(object_repository_path.read_text(encoding="utf-8"))

    steps = read_steps(test_data_path)
    validate_steps(steps)
    test_cases = group_steps(steps)
    driver = build_driver()
    actions = ActionKeywords(driver, options["frontendBaseUrl"], options["apiBaseUrl"], object_repository)
    results = []
    try:
        for test_case in test_cases:
            result = run_case(driver, actions, options["frontendBaseUrl"], test_case)
            results.append(result)
            print(f"{result['status'].upper()} {result['id']} {result['name']}")
            if result["status"] == "failed":
                print(f"FAILED DETAIL {result['id']}: {result['error']}", file=sys.stderr)
    finally:
        driver.quit()

    summary = {
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "framework": "keyword-driven",
        "components": {
            "excel": str(test_data_path),
            "objectRepository": str(object_repository_path),
            "functionLibrary": str(Path(__file__).resolve().parent / "action_keywords.py"),
            "driverScript": str(Path(__file__).resolve()),
            "selenium": "selenium",
        },
        "frontendBaseUrl": options["frontendBaseUrl"],
        "apiBaseUrl": options["apiBaseUrl"],
        "total": len(results),
        "passed": len([item for item in results if item["status"] == "passed"]),
        "failed": len([item for item in results if item["status"] == "failed"]),
        "results": results,
    }
    write_json_report(Path(options["reportPath"]), summary)
    write_html_report(
        Path(options["htmlReportPath"]),
        "Dahla Keyword-Driven Website Report",
        "Website automation scenarios loaded from Excel TestSteps and executed through Python ActionKeywords.",
        "keyword-driven",
        summary,
    )
    return summary
