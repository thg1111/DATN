"""
Script tạo file Excel test data cho AutomationFramework.
Chạy 1 lần: python _generate_excel.py
Sinh ra: test_data.xlsx (3 sheet: Login, Register, Checkout)
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="9F2F49", end_color="9F2F49", fill_type="solid")
DATA_FONT = Font(name="Arial", size=10)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
PASS_FILL = PatternFill(start_color="E8F5EC", end_color="E8F5EC", fill_type="solid")
FAIL_FILL = PatternFill(start_color="FDEBEA", end_color="FDEBEA", fill_type="solid")


def style_header(ws, row_num, col_count):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER


def style_data(ws, row_num, col_count):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = DATA_FONT
        cell.alignment = ALIGN_LEFT
        cell.border = THIN_BORDER


def create_login_sheet(wb):
    ws = wb.create_sheet("Login")
    headers = ["id", "name", "username", "password", "expected"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    style_header(ws, 1, len(headers))

    data = [
        ["L01", "Admin đăng nhập thành công", "admin", "123", "success"],
        ["L02", "Bỏ trống username", "", "123", "html5-invalid"],
        ["L03", "Bỏ trống password", "admin", "", "html5-invalid"],
        ["L04", "Sai mật khẩu", "admin", "wrongpassword", "success"],
        ["L05", "Username không tồn tại", "nouser", "123", "success"],
        ["L06", "Bỏ trống cả 2 trường", "", "", "html5-invalid"],
        ["L07", "Username có khoảng trắng", "  admin  ", "123", "success"],
        ["L08", "Mật khẩu có ký tự đặc biệt", "admin", "P@ss!123#", "success"],
    ]
    for r_idx, row_data in enumerate(data, 2):
        for c_idx, val in enumerate(row_data, 1):
            ws.cell(row=r_idx, column=c_idx, value=val)
        style_data(ws, r_idx, len(headers))

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 25
    ws.column_dimensions["D"].width = 25
    ws.column_dimensions["E"].width = 18
    return ws


def create_register_sheet(wb):
    ws = wb.create_sheet("Register")
    headers = ["id", "name", "tenTK", "matKhau", "email", "tenND", "sdt", "sinhNhat", "diaChi", "gioiTinh", "expected"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    style_header(ws, 1, len(headers))

    data = [
        ["R01", "Đăng ký hợp lệ đầy đủ", "testuser01", "Pass123!", "test01@mail.com", "Nguyễn Văn A", "0123456789", "2000-01-15", "Hưng Yên", "Nam", "valid-form"],
        ["R02", "Thiếu email - invalid", "testuser02", "Pass123!", "", "Nguyễn Văn B", "0123456789", "", "", "Nam", "html5-invalid"],
        ["R03", "Thiếu username - invalid", "", "Pass123!", "test03@mail.com", "Nguyễn Văn C", "0123456789", "", "", "Nữ", "html5-invalid"],
        ["R04", "Thiếu mật khẩu - invalid", "testuser04", "", "test04@mail.com", "Nguyễn Văn D", "0123456789", "", "", "Nam", "html5-invalid"],
        ["R05", "Thiếu họ tên - invalid", "testuser05", "Pass123!", "test05@mail.com", "", "0123456789", "", "", "Nam", "html5-invalid"],
        ["R06", "Thiếu SĐT - invalid", "testuser06", "Pass123!", "test06@mail.com", "Nguyễn Văn F", "", "", "", "Nữ", "html5-invalid"],
        ["R07", "Đăng ký nữ hợp lệ", "testuser07", "SecureP@ss", "nu07@mail.com", "Trần Thị G", "0987654321", "1995-06-20", "Hà Nội", "Nữ", "valid-form"],
        ["R08", "Đầy đủ thông tin với địa chỉ", "testuser08", "MyP@ss99", "full08@mail.com", "Lê Văn H", "0369852147", "1998-12-01", "Khoái Châu, Hưng Yên", "Nam", "valid-form"],
    ]
    for r_idx, row_data in enumerate(data, 2):
        for c_idx, val in enumerate(row_data, 1):
            ws.cell(row=r_idx, column=c_idx, value=val)
        style_data(ws, r_idx, len(headers))

    for col_letter, w in [("A", 8), ("B", 30), ("C", 18), ("D", 16), ("E", 22), ("F", 20), ("G", 14), ("H", 14), ("I", 25), ("J", 12), ("K", 16)]:
        ws.column_dimensions[col_letter].width = w
    return ws


def create_checkout_sheet(wb):
    ws = wb.create_sheet("Checkout")
    headers = [
        "id", "name",
        "productId", "productName", "productPrice", "productQuantity",
        "loginMaND", "loginTenTK",
        "customerName", "customerPhone", "customerEmail",
        "customerAddress", "customerGender", "customerMessage",
        "expected",
    ]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    style_header(ws, 1, len(headers))

    data = [
        ["C01", "Giỏ hàng trống - hiển thị thông báo",
         "", "", "", "",
         "", "",
         "", "", "", "", "", "",
         "cart-empty"],
        ["C02", "Thêm 1 sản phẩm vào giỏ hàng",
         "101", "Hoa hồng đỏ", 350000, 1,
         "", "",
         "", "", "", "", "", "",
         "cart-has-items"],
        ["C03", "Thêm sản phẩm số lượng lớn (x10)",
         "101", "Hoa hồng đỏ", 350000, 10,
         "", "",
         "", "", "", "", "", "",
         "cart-has-items"],
        ["C04", "Thêm sản phẩm giá cao",
         "2", "Pure Enchanter", 11450000, 1,
         "", "",
         "", "", "", "", "", "",
         "cart-has-items"],
        ["C05", "Đặt hàng khi chưa đăng nhập - bị chặn",
         "101", "Hoa hồng đỏ", 350000, 1,
         "", "",
         "Nguyễn Văn Test", "0364239807", "test@mail.com", "Hưng Yên", "male", "",
         "invoice-blocked-no-login"],
        ["C06", "Đặt hàng khi giỏ trống - bị chặn",
         "", "", "", "",
         1, "admin",
         "Admin Test", "0364239807", "", "", "male", "",
         "invoice-blocked-empty-cart"],
        ["C07", "Đặt hàng thành công - đầy đủ thông tin",
         "101", "Hoa hồng đỏ", 350000, 2,
         1, "admin",
         "Nguyễn Thị Hoa", "0364239807", "hoa@dahla.vn", "Khoái Châu, Hưng Yên", "female", "Giao trước 17h",
         "invoice-success"],
        ["C08", "Đặt hàng thành công - tối thiểu thông tin",
         "23", "Kiêu sa", 4600000, 1,
         1, "admin",
         "Trần Văn Nam", "0987654321", "", "", "male", "",
         "invoice-success"],
        ["C09", "Đặt hàng sản phẩm giáng sinh",
         "41", "Phút giây an lành", 540000, 3,
         1, "admin",
         "Lê Thị Mai", "0912345678", "mai@mail.com", "TP. Hưng Yên", "female", "Gói quà đẹp nhé",
         "invoice-success"],
        ["C10", "Đặt hàng gấu bông",
         "33", "Teddy", 450000, 1,
         1, "admin",
         "Phạm Quốc Việt", "0368741256", "", "Yên Mỹ, Hưng Yên", "male", "",
         "invoice-success"],
    ]
    for r_idx, row_data in enumerate(data, 2):
        for c_idx, val in enumerate(row_data, 1):
            ws.cell(row=r_idx, column=c_idx, value=val)
        style_data(ws, r_idx, len(headers))

    widths = [8, 40, 10, 22, 14, 14, 12, 14, 22, 16, 22, 28, 14, 24, 26]
    for i, w in enumerate(widths):
        col_letter = chr(65 + i) if i < 26 else chr(64 + i // 26) + chr(65 + i % 26)
        ws.column_dimensions[col_letter].width = w
    return ws


def main():
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    create_login_sheet(wb)
    create_register_sheet(wb)
    create_checkout_sheet(wb)

    out_path = Path(__file__).parent / "test_data.xlsx"
    wb.save(str(out_path))
    print(f"[OK] Da tao: {out_path}")
    print(f"   - Sheet Login: 8 test cases")
    print(f"   - Sheet Register: 8 test cases")
    print(f"   - Sheet Checkout: 10 test cases")


if __name__ == "__main__":
    main()
