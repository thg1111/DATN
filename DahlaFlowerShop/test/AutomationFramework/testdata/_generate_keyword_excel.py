"""
Script tạo file Excel keyword steps cho Keyword-Driven Testing.
Sinh ra: keyword_steps.xlsx (sheet TestSteps)
Bao gồm: TC trang chủ, đăng nhập, thêm giỏ hàng, đặt hàng.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
DATA_FONT = Font(name="Arial", size=10)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def style_header(ws, row_num, col_count):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER


def style_data(ws, row_num, col_count):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = DATA_FONT
        cell.alignment = ALIGN_LEFT
        cell.border = THIN_BORDER


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "TestSteps"

    headers = ["testCaseId", "testCaseName", "stepOrder", "keyword", "objectName", "testData", "description", "runFlag"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    style_header(ws, 1, len(headers))

    steps = [
        # TC01: Trang chu hien thi
        ["TC01", "Trang chu hien thi dung", 1, "open", "homePage", "", "Mo trang chu", "Y"],
        ["TC01", "Trang chu hien thi dung", 2, "configureApi", "", "", "Cau hinh API URL", "Y"],
        ["TC01", "Trang chu hien thi dung", 3, "waitVisible", "siteHeader", "", "Cho header hien", "Y"],
        ["TC01", "Trang chu hien thi dung", 4, "waitVisible", "siteMenu", "", "Cho menu hien", "Y"],

        # TC02: Them san pham vao gio hang
        ["TC02", "Them san pham vao gio hang", 1, "open", "categoryPage", "", "Mo trang danh muc", "Y"],
        ["TC02", "Them san pham vao gio hang", 2, "configureApi", "", "", "Cau hinh API", "Y"],
        ["TC02", "Them san pham vao gio hang", 3, "click", "categoryAddToCartButton", "", "Click them gio hang", "Y"],
        ["TC02", "Them san pham vao gio hang", 4, "assertLocalStorageObjectNotEmpty", "cartStorage", "", "Kiem tra gio hang co du lieu", "Y"],

        # TC03: Dang nhap thanh cong
        ["TC03", "Dang nhap thanh cong", 1, "open", "loginPage", "", "Mo trang dang nhap", "Y"],
        ["TC03", "Dang nhap thanh cong", 2, "configureApi", "", "", "Cau hinh API", "Y"],
        ["TC03", "Dang nhap thanh cong", 3, "type", "loginUsernameInput", "admin", "Nhap username", "Y"],
        ["TC03", "Dang nhap thanh cong", 4, "type", "loginPasswordInput", "123", "Nhap password", "Y"],
        ["TC03", "Dang nhap thanh cong", 5, "click", "loginSubmitButton", "", "Click dang nhap", "Y"],

        # TC04: Mo trang gio hang
        ["TC04", "Mo trang gio hang", 1, "open", "cartPage", "", "Mo trang gio hang", "Y"],
        ["TC04", "Mo trang gio hang", 2, "configureApi", "", "", "Cau hinh API", "Y"],
        ["TC04", "Mo trang gio hang", 3, "waitVisible", "cartContent", "", "Cho noi dung gio hang hien", "Y"],

        # TC05: Luong dat hang day du
        ["TC05", "Luong dat hang day du", 1, "open", "loginPage", "", "Mo trang dang nhap", "Y"],
        ["TC05", "Luong dat hang day du", 2, "configureApi", "", "", "Cau hinh API", "Y"],
        ["TC05", "Luong dat hang day du", 3, "type", "loginUsernameInput", "admin", "Nhap username", "Y"],
        ["TC05", "Luong dat hang day du", 4, "type", "loginPasswordInput", "123", "Nhap password", "Y"],
        ["TC05", "Luong dat hang day du", 5, "click", "loginSubmitButton", "", "Click dang nhap", "Y"],
        ["TC05", "Luong dat hang day du", 6, "open", "cartPage", "", "Mo trang gio hang", "Y"],
        ["TC05", "Luong dat hang day du", 7, "configureApi", "", "", "Cau hinh API", "Y"],
        ["TC05", "Luong dat hang day du", 8, "waitVisible", "cartContent", "", "Cho gio hang hien", "Y"],
        ["TC05", "Luong dat hang day du", 9, "type", "checkoutNameInput", "Nguyen Van Test", "Nhap ho ten", "Y"],
        ["TC05", "Luong dat hang day du", 10, "type", "checkoutPhoneInput", "0364239807", "Nhap SDT", "Y"],
        ["TC05", "Luong dat hang day du", 11, "type", "checkoutAddressTextarea", "Khoai Chau, Hung Yen", "Nhap dia chi", "Y"],
        ["TC05", "Luong dat hang day du", 12, "click", "createInvoiceButton", "", "Click tao hoa don", "Y"],
    ]

    for r_idx, row_data in enumerate(steps, 2):
        for c_idx, val in enumerate(row_data, 1):
            ws.cell(row=r_idx, column=c_idx, value=val)
        style_data(ws, r_idx, len(headers))

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 38
    ws.column_dimensions["E"].width = 30
    ws.column_dimensions["F"].width = 25
    ws.column_dimensions["G"].width = 35
    ws.column_dimensions["H"].width = 10

    out_path = Path(__file__).parent / "keyword_steps.xlsx"
    wb.save(str(out_path))
    print(f"[OK] Da tao: {out_path}")
    print(f"   - Sheet TestSteps: {len(steps)} steps, 5 test cases")


if __name__ == "__main__":
    main()
