"""
ExcelReader - Đọc dữ liệu test từ file Excel (.xlsx).
Hỗ trợ đọc các sheet: Login, Register, Checkout, TestSteps.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def clean(value: Any) -> str:
    """Chuẩn hoá giá trị cell: bỏ None, strip khoảng trắng."""
    return str(value or "").strip()


def has_data(row: dict[str, Any]) -> bool:
    """Kiểm tra dòng có dữ liệu hay không."""
    return any(clean(v) for v in row.values())


def read_sheet(workbook, sheet_name: str) -> list[dict[str, Any]]:
    """Đọc 1 sheet thành danh sách dict (header → value)."""
    if sheet_name not in workbook.sheetnames:
        raise RuntimeError(f'Không tìm thấy sheet "{sheet_name}" trong file Excel.')

    sheet = workbook[sheet_name]
    headers = [clean(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows: list[dict[str, Any]] = []
    for excel_row in sheet.iter_rows(min_row=2, values_only=True):
        row = {
            headers[i]: excel_row[i] if i < len(excel_row) else ""
            for i in range(len(headers))
        }
        if has_data(row):
            rows.append(row)
    return rows


def require_columns(rows: list[dict[str, Any]], sheet_name: str, required: list[str]) -> None:
    """Kiểm tra sheet có đủ các cột bắt buộc."""
    available = set()
    for row in rows:
        available.update(row.keys())
    missing = [col for col in required if col not in available]
    if missing:
        raise RuntimeError(f'Sheet "{sheet_name}" thiếu cột: {", ".join(missing)}')


class ExcelReader:
    """Đọc file Excel test data cho AutomationFramework."""

    def __init__(self, file_path: str | Path):
        self.file_path = Path(file_path)
        if not self.file_path.exists():
            raise FileNotFoundError(f"Không tìm thấy file Excel: {self.file_path}")
        self.workbook = load_workbook(str(self.file_path), data_only=True)

    # ──────────── Login ────────────
    def read_login_cases(self) -> list[dict[str, Any]]:
        """Đọc sheet Login → danh sách test case đăng nhập."""
        rows = read_sheet(self.workbook, "Login")
        require_columns(rows, "Login", ["id", "name", "username", "password", "expected"])

        return [
            {
                "id": clean(row.get("id")),
                "name": clean(row.get("name")),
                "username": clean(row.get("username")),
                "password": clean(row.get("password")),
                "expected": clean(row.get("expected")),
            }
            for row in rows
        ]

    # ──────────── Register ────────────
    def read_register_cases(self) -> list[dict[str, Any]]:
        """Đọc sheet Register → danh sách test case đăng ký."""
        rows = read_sheet(self.workbook, "Register")
        require_columns(rows, "Register", ["id", "name", "tenTK", "matKhau", "email", "tenND", "sdt", "expected"])

        return [
            {
                "id": clean(row.get("id")),
                "name": clean(row.get("name")),
                "data": {
                    "tenTK": clean(row.get("tenTK")),
                    "matKhau": clean(row.get("matKhau")),
                    "email": clean(row.get("email")),
                    "tenND": clean(row.get("tenND")),
                    "sdt": clean(row.get("sdt")),
                    "sinhNhat": clean(row.get("sinhNhat")),
                    "diaChi": clean(row.get("diaChi")),
                    "gioiTinh": clean(row.get("gioiTinh")),
                },
                "expected": clean(row.get("expected")),
            }
            for row in rows
        ]

    # ──────────── Checkout ────────────
    def read_checkout_cases(self) -> list[dict[str, Any]]:
        """Đọc sheet Checkout → danh sách test case đặt hàng."""
        rows = read_sheet(self.workbook, "Checkout")
        require_columns(rows, "Checkout", ["id", "name", "expected"])

        cases = []
        for row in rows:
            product_id = clean(row.get("productId"))
            case = {
                "id": clean(row.get("id")),
                "name": clean(row.get("name")),
                "product": {
                    "id": product_id,
                    "name": clean(row.get("productName")),
                    "price": int(row.get("productPrice") or 0),
                    "quantity": int(row.get("productQuantity") or 0),
                } if product_id else None,
                "login": {
                    "maND": int(row.get("loginMaND") or 0),
                    "tenTK": clean(row.get("loginTenTK")),
                } if clean(row.get("loginMaND")) else None,
                "customer": {
                    "fullname": clean(row.get("customerName")),
                    "phone": clean(row.get("customerPhone")),
                    "email": clean(row.get("customerEmail")),
                    "address": clean(row.get("customerAddress")),
                    "gender": clean(row.get("customerGender")) or "male",
                    "message": clean(row.get("customerMessage")),
                } if clean(row.get("customerName")) else None,
                "expected": clean(row.get("expected")),
            }
            cases.append(case)
        return cases

    # ──────────── Keyword Steps ────────────
    def read_keyword_steps(self, sheet_name: str = "TestSteps") -> list[dict[str, Any]]:
        """Đọc sheet TestSteps → danh sách keyword steps (cho keyword-driven)."""
        actual_sheet = sheet_name if sheet_name in self.workbook.sheetnames else "WebsiteKeywords"
        rows = read_sheet(self.workbook, actual_sheet)

        steps = []
        for row in rows:
            run_flag = clean(row.get("runFlag") or "Y").upper()
            if run_flag == "N":
                continue
            steps.append({
                "testCaseId": clean(row.get("testCaseId") or row.get("id")),
                "testCaseName": clean(row.get("testCaseName") or row.get("name")),
                "stepOrder": int(row.get("stepOrder") or 0),
                "keyword": clean(row.get("keyword")),
                "objectName": clean(row.get("objectName") or row.get("target")),
                "testData": clean(row.get("testData") or row.get("value")),
                "description": clean(row.get("description")),
            })
        return steps

    def group_keyword_steps(self, sheet_name: str = "TestSteps") -> list[dict[str, Any]]:
        """Nhóm keyword steps theo testCaseId."""
        steps = self.read_keyword_steps(sheet_name)
        cases: dict[str, dict[str, Any]] = {}
        for step in steps:
            tc = cases.setdefault(
                step["testCaseId"],
                {"id": step["testCaseId"], "name": step["testCaseName"] or step["testCaseId"], "steps": []},
            )
            tc["steps"].append(step)

        return [
            {**tc, "steps": sorted(tc["steps"], key=lambda s: s["stepOrder"])}
            for tc in cases.values()
        ]
