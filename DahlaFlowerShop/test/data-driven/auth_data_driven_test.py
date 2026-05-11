from __future__ import annotations

import os
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urljoin

from openpyxl import load_workbook
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

sys.path.append(str(Path(__file__).resolve().parents[1]))
from report_utils import write_html_report, write_json_report  # noqa: E402


BASE_DIR = Path(__file__).resolve().parent
TEST_DIR = BASE_DIR.parent
REPORTS_DIR = TEST_DIR / "reports"
EXCEL_PATH = BASE_DIR / "auth-test-data-styled.xlsx"
REPORT_JSON = REPORTS_DIR / "data-driven-auth-report.json"
REPORT_HTML = REPORTS_DIR / "data-driven-auth-report.html"

FRONTEND_BASE_URL = os.getenv("FRONTEND_BASE_URL", "http://127.0.0.1:5500")
API_BASE_URL = os.getenv("API_BASE_URL", "https://localhost:7114")

LOGIN_REQUIRED_COLUMNS = ["id", "name", "username", "password", "expected"]
REGISTER_REQUIRED_COLUMNS = ["id", "name", "tenTK", "matKhau", "email", "tenND", "sdt", "expected"]
VALID_LOGIN_EXPECTED = {"success", "html5-invalid"}
VALID_REGISTER_EXPECTED = {"valid-form", "html5-invalid"}


def clean(value: Any) -> str:
    return str(value or "").strip()


def has_data(row: dict[str, Any]) -> bool:
    return any(clean(value) for value in row.values())


def read_sheet(workbook: Any, sheet_name: str) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        raise RuntimeError(f'Missing sheet "{sheet_name}" in {EXCEL_PATH}')

    sheet = workbook[sheet_name]
    headers = [clean(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows: list[dict[str, Any]] = []
    for excel_row in sheet.iter_rows(min_row=2, values_only=True):
        row = {headers[index]: excel_row[index] if index < len(excel_row) else "" for index in range(len(headers))}
        if has_data(row):
            rows.append(row)
    return rows


def require_columns(rows: list[dict[str, Any]], sheet_name: str, required_columns: list[str]) -> None:
    available = set()
    for row in rows:
        available.update(row.keys())
    missing = [column for column in required_columns if column not in available]
    if missing:
        raise RuntimeError(f'Sheet "{sheet_name}" is missing column(s): {", ".join(missing)}')


def validate_expected(test_case: dict[str, Any], sheet_name: str, valid_values: set[str]) -> None:
    if test_case["expected"] not in valid_values:
        raise RuntimeError(
            f'Sheet "{sheet_name}" row "{test_case.get("id") or test_case.get("name")}" '
            f'has invalid expected value "{test_case["expected"]}".'
        )


def read_test_data() -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    workbook = load_workbook(EXCEL_PATH, data_only=True)
    login_rows = read_sheet(workbook, "Login")
    register_rows = read_sheet(workbook, "Register")
    require_columns(login_rows, "Login", LOGIN_REQUIRED_COLUMNS)
    require_columns(register_rows, "Register", REGISTER_REQUIRED_COLUMNS)

    login_cases = [
        {
            "id": clean(row.get("id")),
            "name": clean(row.get("name")),
            "username": clean(row.get("username")),
            "password": clean(row.get("password")),
            "expected": clean(row.get("expected")),
        }
        for row in login_rows
    ]
    register_cases = [
        {
            "id": clean(row.get("id")),
            "name": clean(row.get("name")),
            "data": {
                "tenTK": clean(row.get("tenTK")),
                "matKhau": clean(row.get("matKhau")),
                "email": clean(row.get("email")),
                "tenND": clean(row.get("tenND")),
                "sdt": clean(row.get("sdt")),
                "sinhNhat": clean(row.get("sinhNhat")),
                "diaChi": clean(row.get("diaChi")),
                "gioiTinh": clean(row.get("gioiTinh")),
            },
            "expected": clean(row.get("expected")),
        }
        for row in register_rows
    ]

    for test_case in login_cases:
        validate_expected(test_case, "Login", VALID_LOGIN_EXPECTED)
    for test_case in register_cases:
        validate_expected(test_case, "Register", VALID_REGISTER_EXPECTED)
    return login_cases, register_cases


def page_url(relative_path: str) -> str:
    return urljoin(FRONTEND_BASE_URL.rstrip("/") + "/", relative_path.lstrip("/"))


def build_driver() -> webdriver.Chrome:
    options = Options()
    options.add_argument("--headless=new")
    options.add_argument("--disable-gpu")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--ignore-certificate-errors")
    options.add_argument("--window-size=1366,768")
    return webdriver.Chrome(options=options)


def configure_app(driver: webdriver.Chrome) -> None:
    driver.execute_script(
        """
        window.APP_API_BASE_URL = arguments[0];
        window.alert = function () {};
        window.confirm = function () { return true; };
        """,
        API_BASE_URL,
    )


def clear_state(driver: webdriver.Chrome) -> None:
    driver.execute_script("localStorage.clear(); sessionStorage.clear();")


def wait_visible(driver: webdriver.Chrome, selector: str):
    wait = WebDriverWait(driver, 10)
    return wait.until(EC.visibility_of_element_located((By.CSS_SELECTOR, selector)))


def set_value(driver: webdriver.Chrome, selector: str, value: str) -> None:
    element = wait_visible(driver, selector)
    element.clear()
    if value:
        element.send_keys(value)


def is_form_valid(driver: webdriver.Chrome, selector: str) -> bool:
    return bool(driver.execute_script("return document.querySelector(arguments[0]).checkValidity();", selector))


def run_login_case(driver: webdriver.Chrome, test_case: dict[str, Any]) -> dict[str, Any]:
    driver.get(page_url("/pages/auth/login.html"))
    configure_app(driver)
    clear_state(driver)
    set_value(driver, "#tenTK", test_case["username"])
    set_value(driver, "#matKhau", test_case["password"])

    form_is_valid = is_form_valid(driver, 'form[ng-submit="Login()"]')
    if test_case["expected"] == "html5-invalid":
        if form_is_valid:
            raise RuntimeError("Expected login form to be blocked by HTML5 validation.")
        return {"actual": "html5-invalid"}

    driver.find_element(By.CSS_SELECTOR, 'button[type="submit"]').click()
    WebDriverWait(driver, 15).until(
        lambda current_driver: current_driver.execute_script(
            """
            const session = JSON.parse(localStorage.getItem('DahlaAuthSession') || 'null');
            return Boolean((session && session.tenTK) || localStorage.getItem('MaND'));
            """
        )
    )
    session_name = driver.execute_script(
        """
        const session = JSON.parse(localStorage.getItem('DahlaAuthSession') || 'null');
        return session && session.tenTK;
        """
    )
    return {"actual": "success", "sessionName": session_name}


def run_register_case(driver: webdriver.Chrome, test_case: dict[str, Any]) -> dict[str, Any]:
    driver.get(page_url("/pages/auth/register.html"))
    configure_app(driver)
    clear_state(driver)
    data = test_case["data"]
    set_value(driver, "#tenTK", data["tenTK"])
    set_value(driver, "#matKhau", data["matKhau"])
    set_value(driver, "#email", data["email"])
    set_value(driver, "#tenND", data["tenND"])
    set_value(driver, "#sdt", data["sdt"])
    set_value(driver, "#sinhNhat", data["sinhNhat"])
    set_value(driver, "#diaChi", data["diaChi"])

    if data["gioiTinh"]:
        gender_selector = "#Nam" if data["gioiTinh"] == "Nam" else "#Nu"
        driver.find_element(By.CSS_SELECTOR, gender_selector).click()

    form_is_valid = is_form_valid(driver, "#registerForm")
    if test_case["expected"] == "valid-form":
        if not form_is_valid:
            raise RuntimeError("Expected register form to be valid for this data row.")
        return {"actual": "valid-form"}

    if form_is_valid:
        raise RuntimeError("Expected register form to be blocked by HTML5 validation.")
    return {"actual": "html5-invalid"}


def execute_case(driver: webdriver.Chrome, group: str, test_case: dict[str, Any], test_function) -> dict[str, Any]:
    started_at = time.time()
    try:
        detail = test_function(driver, test_case)
        return {
            "id": test_case["id"],
            "group": group,
            "name": test_case["name"],
            "status": "passed",
            "expected": test_case["expected"],
            "detail": detail,
            "timeMs": int((time.time() - started_at) * 1000),
        }
    except Exception as error:
        return {
            "id": test_case["id"],
            "group": group,
            "name": test_case["name"],
            "status": "failed",
            "expected": test_case["expected"],
            "error": str(error),
            "timeMs": int((time.time() - started_at) * 1000),
        }


def main() -> int:
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    login_cases, register_cases = read_test_data()
    results: list[dict[str, Any]] = []
    driver = build_driver()
    try:
        for test_case in login_cases:
            result = execute_case(driver, "login", test_case, run_login_case)
            results.append(result)
            print(f"{result['status'].upper()} {result['id']} {result['name']}")
        for test_case in register_cases:
            result = execute_case(driver, "register", test_case, run_register_case)
            results.append(result)
            print(f"{result['status'].upper()} {result['id']} {result['name']}")
    finally:
        driver.quit()

    summary = {
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "framework": "data-driven",
        "frontendBaseUrl": FRONTEND_BASE_URL,
        "apiBaseUrl": API_BASE_URL,
        "total": len(results),
        "passed": len([item for item in results if item["status"] == "passed"]),
        "failed": len([item for item in results if item["status"] == "failed"]),
        "results": results,
    }
    write_json_report(REPORT_JSON, summary)
    write_html_report(REPORT_HTML, "Dahla Data-Driven Auth Report", "Login and register test cases loaded from Excel.", "data-driven", summary)
    return 1 if summary["failed"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
